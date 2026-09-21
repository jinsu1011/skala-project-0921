"""Fill the official SKALA report template (.docx) with Markdown content, then convert to PDF.

Supported Markdown subset: #/##/###/#### headings, paragraphs, **bold**, `code`, pipe tables,
'- ' bullets (2-space nesting), '1. ' numbered items, > notes, ```code``` fences, ![caption](image), '---pagebreak---'.
"""
from __future__ import annotations

import copy
import re
import shutil
import subprocess
from dataclasses import dataclass, field
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "report" / "template" / "SKALA_report_template.docx"
NAVY, PURPLE, HEAD_FILL, GRID = "161A58", "7F4ACB", "F2F3F8", "D8DAE5"
KO_FONT = "맑은 고딕"
PDF_FONT_FALLBACK = "Apple SD Gothic Neo"  # used for PDF conversion on machines without Malgun Gothic


@dataclass
class CoverInfo:
    title: str
    cohort: str = "4"
    track: str = "AI"
    team_label: str = "판교 9반 1조"
    class_team: str = "9반 1조"
    members: list[tuple[str, str]] = field(default_factory=list)
    date: str = ""
    instructor: str = "배기주"
    version: str = "v1.0"
    report_kind: str = "과제 제출 보고서"


# ------------------------------------------------------------------ low-level helpers
def _set_run_font(run, size=None, bold=None, color=None, mono=False):
    font = "D2Coding" if mono else KO_FONT
    run.font.name = font
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(a), font if not mono or a != "w:eastAsia" else KO_FONT)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


ZWSP = "\u200b"


def _add_inline(par, text, size=10, color=None, base_bold=False, in_table=False):
    """**bold** and `code` inline markup. In table cells, code gets break points after '_' and ','
    so long identifiers wrap at word parts instead of mid-word."""
    for tok in re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text):
        if not tok:
            continue
        if tok.startswith("**"):
            _set_run_font(par.add_run(tok[2:-2].replace("`", "")), size, True, color or NAVY)
        elif tok.startswith("`"):
            code = tok[1:-1]
            if in_table:
                code = code.replace("_", "_" + ZWSP).replace(",", "," + ZWSP)
            r = par.add_run(code)
            _set_run_font(r, size - (1.0 if in_table else 0.5), base_bold, "5B3FA0", mono=True)
        else:
            _set_run_font(par.add_run(tok), size, base_bold, color)


def _shade(cell, fill):
    tcpr = cell._element.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcpr.append(shd)


def _cell_margins(cell, tb=70, lr=90):
    tcpr = cell._element.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for side, v in (("top", tb), ("left", lr), ("bottom", tb), ("right", lr)):
        e = OxmlElement(f"w:{side}")
        e.set(qn("w:w"), str(v))
        e.set(qn("w:type"), "dxa")
        mar.append(e)
    tcpr.append(mar)


def _table_borders(table):
    tblpr = table._element.tblPr
    b = OxmlElement("w:tblBorders")
    for side, val, color, sz in (("top", "single", NAVY, 8), ("left", "none", "FFFFFF", 0),
                                 ("bottom", "single", NAVY, 8), ("right", "none", "FFFFFF", 0),
                                 ("insideH", "single", GRID, 2), ("insideV", "single", GRID, 2)):
        e = OxmlElement(f"w:{side}")
        e.set(qn("w:val"), val)
        e.set(qn("w:color"), color)
        e.set(qn("w:sz"), str(sz))
        b.append(e)
    tblpr.append(b)


def _fix_grid(table, widths_cm):
    """Fixed layout + explicit grid so LibreOffice/Word honour the column widths."""
    tbl = table._tbl
    tblpr = tbl.tblPr
    for tag in ("w:tblW", "w:tblLayout"):
        old = tblpr.find(qn(tag))
        if old is not None:
            tblpr.remove(old)
    tw = OxmlElement("w:tblW")
    tw.set(qn("w:w"), str(int(sum(widths_cm) * 567)))
    tw.set(qn("w:type"), "dxa")
    tblpr.append(tw)
    lay = OxmlElement("w:tblLayout")
    lay.set(qn("w:type"), "fixed")
    tblpr.append(lay)
    grid = tbl.tblGrid
    for gc, w in zip(grid.findall(qn("w:gridCol")), widths_cm):
        gc.set(qn("w:w"), str(int(w * 567)))


def _repeat_header(row):
    trpr = row._tr.get_or_add_trPr()
    e = OxmlElement("w:tblHeader")
    e.set(qn("w:val"), "true")
    trpr.append(e)


def _no_split(row):
    trpr = row._tr.get_or_add_trPr()
    trpr.append(OxmlElement("w:cantSplit"))


# ------------------------------------------------------------------ builder
class ReportBuilder:
    def __init__(self, cover: CoverInfo):
        self.doc = Document(str(TEMPLATE))
        self.cover = cover
        self.body = self.doc.element.body
        self._fill_cover_and_info()
        self._cut_template_body()

    # ---- template placeholders
    def _replace_everywhere(self, mapping: dict[str, str]):
        parts = [self.doc.element] + [s.header._element for s in self.doc.sections] + \
                [s.first_page_header._element for s in self.doc.sections] + [s.footer._element for s in self.doc.sections]
        for part in parts:
            for t in part.iter(qn("w:t")):
                for k, v in mapping.items():
                    if t.text and k in t.text:
                        t.text = t.text.replace(k, v)

    def _fill_cover_and_info(self):
        c = self.cover
        # cover member lines: template has 1 x 팀장 + 3 x 팀원 -> 5 x 팀원 (no team lead)
        member_ps = [p for p in self.doc.paragraphs if "[ 학번 ] [ 이름 ]" in p.text]
        while len(member_ps) < len(c.members):
            new = copy.deepcopy(member_ps[-1]._element)
            member_ps[-1]._element.addnext(new)
            member_ps = [p for p in self.doc.paragraphs if "[ 학번 ] [ 이름 ]" in p.text]
        for p, (sid, name) in zip(member_ps, c.members):
            for r in p.runs:
                if r.text.strip() == "팀 장":
                    r.text = "팀 원"
                r.text = r.text.replace("[ 학번 ] [ 이름 ]", f"{sid}  {name}")
        self._replace_everywhere({
            "SKALA [ N ]기 과제 제출 보고서": f"SKALA {c.cohort}기 {c.report_kind}",
            "[ 과제명을 입력하세요 ]": c.title,
            "[ 팀명 / 반 · 조 ]": c.team_label,
            "SKALA [ N ]기 · [ 교육 트랙 ]": f"SKALA {c.cohort}기 · {c.track}",
            "[ N ]반 [ N ]조": c.class_team,
            "[ YYYY. MM. DD. ]": c.date,
            "[ 과제명 ]": c.title,
            "[ 개인 / 팀 ]": "팀",
            "SKALA [ N ]기": f"SKALA {c.cohort}기",
            "[ 이름 ]": ", ".join(n for _, n in c.members),
            "[ 강사명 ]": f"{c.instructor} 교수님",
            "v1.0": c.version,
        })
    def _cut_template_body(self):
        """Remove the template's sample chapters (everything after the document-info table)."""
        children = list(self.body)
        tbl_idx = max(i for i, el in enumerate(children) if el.tag == qn("w:tbl"))
        for el in children[tbl_idx + 1:]:
            if el.tag != qn("w:sectPr"):
                self.body.remove(el)
        self._anchor = self.body.find(qn("w:sectPr"))

    def _append(self, el):
        self._anchor.addprevious(el)

    def _style(self, style_id: str):
        return next(st for st in self.doc.styles if st.style_id == style_id)

    def _new_par(self, style=None):
        p = self.doc.add_paragraph()
        if style:
            p.style = self._style(style)
        self._append(p._element)  # add_paragraph appends before sectPr already; this keeps order explicit
        return p

    # ---- block writers
    def heading(self, text, level):
        p = self._new_par(f"Heading{min(level, 3)}")
        _set_run_font(p.add_run(text), {1: 15, 2: 12.5, 3: 11, 4: 10.5}[level], True, NAVY)
        if level == 1 and not text.startswith("문서"):
            p.paragraph_format.page_break_before = False
        p.paragraph_format.keep_with_next = True
        return p

    def paragraph(self, text, size=10, indent=0.0, color=None, italic=False, align=None):
        p = self._new_par()
        pf = p.paragraph_format
        pf.space_after = Pt(4)
        pf.line_spacing = 1.35
        if indent:
            pf.left_indent = Cm(indent)
        if align == "center":
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _add_inline(p, text, size, color)
        if italic:
            for r in p.runs:
                r.font.italic = True
        self._last_par = p
        return p

    def _bind_lead_in(self):
        """A short lead-in paragraph right before a table/code block stays on the same page as the block."""
        p = getattr(self, "_last_par", None)
        if p is not None and p._element.getnext() is self._anchor and len(p.text) < 250:
            p.paragraph_format.keep_with_next = True

    def bullet(self, text, level=0, numbered: str | None = None):
        p = self._new_par()
        pf = p.paragraph_format
        pf.left_indent = Cm(0.55 + 0.6 * level)
        pf.first_line_indent = Cm(-0.4)
        pf.space_after = Pt(2)
        pf.line_spacing = 1.3
        mark = numbered or ("•" if level == 0 else "–")
        _set_run_font(p.add_run(f"{mark} "), 10, level == 0 and not numbered, PURPLE)
        _add_inline(p, text, 10)
        return p

    def note(self, text):
        p = self.paragraph(text, size=9, color="555A6E", indent=0.2)
        pPr = p._element.get_or_add_pPr()
        bdr = OxmlElement("w:pBdr")
        left = OxmlElement("w:left")
        for k, v in (("w:val", "single"), ("w:sz", "18"), ("w:space", "6"), ("w:color", PURPLE)):
            left.set(qn(k), v)
        bdr.append(left)
        pPr.append(bdr)
        return p

    def code(self, text):
        self._bind_lead_in()
        t = self.doc.add_table(rows=1, cols=1)
        self._append(t._element)
        _fix_grid(t, [16.0])
        _no_split(t.rows[0])
        cell = t.rows[0].cells[0]
        cell.width = Cm(16.0)
        _shade(cell, "F7F7FB")
        _cell_margins(cell, 80, 120)
        cell.paragraphs[0].text = ""
        for i, line in enumerate(text.rstrip("\n").split("\n")):
            p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            _set_run_font(p.add_run(line or " "), 7.5, False, "2B2F4A", mono=True)
        self.paragraph("", size=4)

    def table(self, rows: list[list[str]], widths: list[float] | None = None, font=8.8):
        header, body = rows[0], rows[1:]
        self._bind_lead_in()
        t = self.doc.add_table(rows=len(rows), cols=len(header))
        self._append(t._element)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        _table_borders(t)
        total = 16.0
        if not widths:
            lens = [max(len(r[i]) if i < len(r) else 0 for r in rows) for i in range(len(header))]
            lens = [min(max(l, 4), 60) for l in lens]
            widths = [total * l / sum(lens) for l in lens]
        _fix_grid(t, widths)
        for ri, row in enumerate(rows):
            _no_split(t.rows[ri])
            for ci in range(len(header)):
                cell = t.rows[ri].cells[ci]
                cell.width = Cm(widths[ci])
                _cell_margins(cell)
                txt = row[ci] if ci < len(row) else ""
                p = cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.15
                for k, part in enumerate(txt.split("<br>")):
                    if k:
                        p = cell.add_paragraph()
                        p.paragraph_format.space_after = Pt(0)
                        p.paragraph_format.line_spacing = 1.15
                    if ri == 0:
                        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        _set_run_font(p.add_run(part.replace("**", "").replace("`", "")), font, True, NAVY)
                    else:
                        _add_inline(p, part, font, in_table=True)
                if ri == 0:
                    _shade(cell, HEAD_FILL)
        _repeat_header(t.rows[0])
        # short tables never split across pages; long tables keep the header + first 3 rows together
        for row in (t.rows[:-1] if len(rows) <= 12 else t.rows[:4]):
            for cell in row.cells:
                for par in cell.paragraphs:
                    par.paragraph_format.keep_with_next = True
        self.paragraph("", size=4)

    def image(self, path: Path, caption: str, width_cm=15.5):
        p = self._new_par()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(path), width=Cm(width_cm))
        p.paragraph_format.keep_with_next = True  # figure and caption stay on one page
        self.paragraph(caption, size=8.5, color="555A6E", align="center")

    def page_break(self):
        p = self._new_par()
        p.add_run().add_break(WD_BREAK.PAGE)

    # ---- markdown
    def markdown(self, md: str, base_dir: Path):
        lines = md.split("\n")
        i = 0
        para: list[str] = []
        pending_widths = None

        def flush_para():
            if para:
                text = " ".join(s.strip() for s in para)
                p = self.paragraph(text)
                if text.startswith("**") and text.endswith("**"):  # run-in caption/label -> keep with next block
                    p.paragraph_format.keep_with_next = True
                para.clear()

        while i < len(lines):
            ln = lines[i]
            s = ln.strip()
            if s.startswith("```"):
                flush_para()
                j = i + 1
                while not lines[j].strip().startswith("```"):
                    j += 1
                self.code("\n".join(lines[i + 1:j]))
                i = j + 1
                continue
            if s == "---pagebreak---":
                flush_para()
                self.page_break()
            elif m := re.match(r"^(#{1,4})\s+(.*)", s):
                flush_para()
                self.heading(m.group(2), len(m.group(1)))
            elif s.startswith("<!--w:"):
                flush_para()
                pending_widths = [float(x) for x in s[6:-3].split(",")]
            elif s.startswith("|"):
                flush_para()
                rows, widths = [], pending_widths
                pending_widths = None
                while i < len(lines) and lines[i].strip().startswith("|"):
                    cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                    if all(re.fullmatch(r":?-{2,}:?", c) for c in cells):
                        pass
                    else:
                        rows.append(cells)
                    i += 1
                self.table(rows, widths)
                continue
            elif m := re.match(r"^!\[(.*)\]\((.*?)\)(?:<!--img:(\d+(?:\.\d+)?)-->)?", s):
                flush_para()
                self.image(base_dir / m.group(2), m.group(1), float(m.group(3) or 15.5))
            elif m := re.match(r"^(\s*)- (.*)", ln):
                flush_para()
                b = self.bullet(m.group(2), len(m.group(1)) // 2)
                nxt = lines[i + 1] if i + 1 < len(lines) else ""
                if re.match(r"^\s*- ", nxt) and len(nxt) - len(nxt.lstrip()) > len(m.group(1)):
                    b.paragraph_format.keep_with_next = True
            elif m := re.match(r"^(\s*)(\d+)\. (.*)", ln):
                flush_para()
                self.bullet(m.group(3), len(m.group(1)) // 2, numbered=f"{m.group(2)}.")
            elif s.startswith("> "):
                flush_para()
                self.note(s[2:])
            elif not s:
                flush_para()
            elif s.startswith("<!--"):
                pass
            else:
                para.append(s)
            i += 1
        flush_para()

    def save(self, path: Path):
        path.parent.mkdir(parents=True, exist_ok=True)
        self.doc.save(str(path))
        return path


def _soffice() -> str | None:
    for c in ("soffice", "/Applications/LibreOffice.app/Contents/MacOS/soffice",
              r"C:\Program Files\LibreOffice\program\soffice.exe", "libreoffice"):
        if shutil.which(c) or Path(c).exists():
            return shutil.which(c) or c
    return None


def docx_to_pdf(docx_path: Path, pdf_path: Path) -> Path:
    """LibreOffice headless conversion. Malgun Gothic is swapped for an installed Korean font if missing."""
    exe = _soffice()
    if exe is None:
        raise RuntimeError("LibreOffice (soffice) not found; install it or open the .docx and export to PDF")
    work = pdf_path.parent / f".build_{docx_path.stem}"
    work.mkdir(parents=True, exist_ok=True)
    src = work / "doc.docx"
    has_malgun = subprocess.run(["fc-list", ":family=Malgun Gothic"], capture_output=True, text=True).stdout.strip() \
        if shutil.which("fc-list") else ""
    if not has_malgun and Path("/System/Library/Fonts/AppleSDGothicNeo.ttc").exists():
        _swap_font(docx_path, src, KO_FONT, PDF_FONT_FALLBACK)
    else:
        shutil.copy(docx_path, src)
    subprocess.run([exe, "--headless", "--convert-to", "pdf", "--outdir", str(work), str(src)],
                   check=True, capture_output=True, timeout=240)
    pdf_path.parent.mkdir(parents=True, exist_ok=True)
    shutil.move(str(work / "doc.pdf"), pdf_path)
    shutil.rmtree(work, ignore_errors=True)
    return pdf_path


def _swap_font(src: Path, dst: Path, old: str, new: str):
    import zipfile

    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename.endswith(".xml"):
                data = data.decode("utf8").replace(f'"{old}"', f'"{new}"').encode("utf8")
            zout.writestr(item, data)
